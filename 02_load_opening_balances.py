import pandas as pd
import numpy as np
import openpyxl as op


def handle(file, sheet):
    df1 = pd.read_excel(file, sheet_name=sheet)
    # print(df1.info())
    group = df1.groupby('Code')
    list1 = []
    for key, value in group:
        # print(key)
        df3 = pd.DataFrame(value)
        # print(df3)
        list1.append(df3.index[0])
    # print(list1)
    list_code = []
    for i in list1:
        # a = '0' * (6 - len(str(int(df1.iloc[i]['Code'])))) + str(int(df1.iloc[i]['Code']))
        a = int(df1.iloc[i]['Code'])
        list_code.append(a)
    # print(list_code)
    dict2 = {}
    for idx, value in enumerate(list_code):
        dict2[value] = list1[idx]
    # print(dict2)

    df2 = pd.read_excel(open_file, header=None)
    df2.columns = ['A', 'B', "C"]
    a = np.arange(65, 91)
    list_ABC = [chr(i) for i in a]
    dict1 = {}
    for i in zip(range(0, 27), list_ABC):
        dict1[i[1]] = i[0]
    # print(df2.iloc[10:,3].columns)
    # print(df3.shape)
    df3 = df2
    df3.columns = list_ABC[0:df3.shape[1]]
    df3['A'] = df3['A'].astype(int)
    df3.dropna(axis=0, subset=['A'])
    # print(df3.info())
    list_null = []
    list_not_null = []
    list_op = []
    dict3 = {}
    for i in list_code:
        # if len(df3[df3['A'].str.contains(i)]['A']) != 0:
        #     a = df3[df3['A'].str.contains(i)][['B', 'C']].values[0]
        #     if a[1] != 0 and a[0] != 0:
        #         list_not_null.append(i)
        #         list_op.append([a[0], a[1], a[1] / a[0]])
        # else:
        #     list_null.append(i)
        if len(df3[df3['A'] == i]['A']) != 0:
            print(i)
            a = df3[df3['A'] == i][['B', 'C']].values[0]
            if a[1] != 0 and a[0] != 0:
                list_not_null.append(i)
                list_op.append([a[0], a[1], a[1] / a[0]])
        else:
            list_null.append(i)
    for idx, index in enumerate(list_not_null):
        dict3[index] = list(list_op[idx])
    # 这一步开始运算了
    # print(dict3)
    table = op.load_workbook(file)
    table1 = table[sheet]
    list_not_null_index = [dict2[i] for i in list_not_null]
    list_null_index = [dict2[i] for i in list_null]
    for idx, i in enumerate(list_not_null_index):
        for j in range(4, 7):
            table1.cell(i + 2, j, dict3[list_not_null[idx]][j - 4])
    for idx, i in enumerate(list_null_index):
        for j in range(4, 7):
            table1.cell(i + 2, j, 0)
    table.save(file)


if __name__ == "__main__":

    open_file = 'input/opening.xlsx' # 这里放opening.xlsx的地址
    file = 'output/trans_result_WA.xlsx' # 这里放trans_result_WA.xlsx的地址
    for i in range(10):
        sheetname = 'transaction' + str(i + 1)
        print(sheetname)
        handle(file, sheetname)